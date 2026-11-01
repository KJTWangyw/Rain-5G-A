#!/usr/bin/env python3
"""Build paper/Manuscript.docx from paper/content.md on the style-only template, then convert it to PDF.

content.md syntax
  ::title / ::abstract / ::keywords   block until a blank line
  # Heading            numbered level-1 heading;  #* Heading  unnumbered;  ## 2.1 Heading  level 2
  $$ LaTeX $$ {#eq:label}            display equation (numbered)
  ![caption](file.png){#fig:label width=14.5}
  ::table tab:label | caption | widths in cm      rows "a | b | c", first row is the header, end with ::end
  ::refs ... ::end                    one reference per line
  @eq:label, @fig:label, @tab:label   replaced by numbers; {{key}} replaced from results/summary.json
Inline: $...$ (small LaTeX subset), *italic*, **bold**.

Usage:  python paper/build_manuscript.py [--no-pdf]
"""
from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
from pathlib import Path

import docx
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))
import docxkit as K  # noqa: E402


def fill_placeholders(text, values):
    missing = sorted(set(re.findall(r"\{\{(\w+)\}\}", text)) - set(values))
    if missing:
        raise KeyError(f"placeholders without values: {missing}")
    return re.sub(r"\{\{(\w+)\}\}", lambda m: str(values[m.group(1)]), text)


def parse(text):
    """Return a list of (kind, payload) blocks."""
    blocks, lines, i = [], text.splitlines(), 0
    para = []

    def flush():
        if para:
            blocks.append(("para", " ".join(s.strip() for s in para)))
            para.clear()

    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("%"):
            i += 1
            continue
        if not line.strip():
            flush()
            i += 1
            continue
        if line.startswith("::title") or line.startswith("::abstract") or line.startswith("::keywords"):
            flush()
            kind = line[2:].strip()
            body = []
            i += 1
            while i < len(lines) and lines[i].strip():
                body.append(lines[i].strip())
                i += 1
            blocks.append((kind, " ".join(body)))
            continue
        if line.startswith("::table"):
            flush()
            head = [h.strip() for h in line[len("::table"):].split("|")]
            rows = []
            i += 1
            while not lines[i].startswith("::end"):
                if lines[i].strip():
                    rows.append([c.strip() for c in lines[i].split("|")])
                i += 1
            blocks.append(("table", {"label": head[0], "caption": head[1],
                                     "widths": [float(w) for w in head[2].split(",")] if len(head) > 2 else None,
                                     "rows": rows}))
            i += 1
            continue
        if line.startswith("::refs"):
            flush()
            refs = []
            i += 1
            while not lines[i].startswith("::end"):
                if lines[i].strip():
                    refs.append(lines[i].strip())
                i += 1
            blocks.append(("refs", refs))
            i += 1
            continue
        m = re.match(r"^\$\$(.+)\$\$\s*\{#(eq:[\w-]+)\}\s*$", line)
        if m:
            flush()
            blocks.append(("equation", {"tex": m.group(1).strip(), "label": m.group(2)}))
            i += 1
            continue
        m = re.match(r"^!\[(.*)\]\((.+?)\)\{#(fig:[\w-]+)\s+width=([\d.]+)\}\s*$", line)
        if m:
            flush()
            blocks.append(("figure", {"caption": m.group(1), "file": m.group(2), "label": m.group(3),
                                      "width": float(m.group(4))}))
            i += 1
            continue
        if line.startswith("#"):
            flush()
            m = re.match(r"^(#+)(\*?)\s+(.*)$", line)
            blocks.append(("heading", {"level": len(m.group(1)), "numbered": not m.group(2), "text": m.group(3)}))
            i += 1
            continue
        para.append(line)
        i += 1
    flush()
    return blocks


def number_labels(blocks):
    numbers, counters = {}, {"eq": 0, "fig": 0, "tab": 0}
    for kind, payload in blocks:
        if kind in ("equation", "figure", "table"):
            prefix = payload["label"].split(":")[0]
            counters[prefix] += 1
            numbers[payload["label"]] = counters[prefix]
    return numbers


def resolve_refs(text, numbers):
    def sub(m):
        if m.group(1) not in numbers:
            raise KeyError(f"unknown label {m.group(1)}")
        return str(numbers[m.group(1)])
    return re.sub(r"@((?:eq|fig|tab):[\w-]+)", sub, text)


def build(content_path, out_docx, summary_path):
    values = json.loads(Path(summary_path).read_text())
    text = fill_placeholders(Path(content_path).read_text(encoding="utf-8"), values)
    blocks = parse(text)
    numbers = number_labels(blocks)
    equations = [p["tex"] for k, p in blocks if k == "equation"]
    images = K.render_equations(equations, HERE / "build" / "equations")
    eq_iter = iter(images)

    doc = docx.Document(str(HERE / "template" / "reference.docx"))
    previous = None
    for kind, payload in blocks:
        if kind == "title":
            p = K.title(doc, resolve_refs(payload, numbers))
            for r in p.runs:
                r.font.size = Pt(12)
        elif kind == "abstract":
            p = K.body_paragraph(doc, indent=False)
            K.add_run(p, "Abstract: ", bold=True)
            K.add_rich_text(p, resolve_refs(payload, numbers))
        elif kind == "keywords":
            p = K.body_paragraph(doc, indent=False)
            K.add_run(p, "Keywords: ", bold=True)
            K.add_rich_text(p, payload)
        elif kind == "heading":
            K.heading(doc, resolve_refs(payload["text"], numbers), payload["level"], payload["numbered"])
        elif kind == "para":
            # a sentence continuing after a display equation ("where ...") is not indented, as in the template
            p = K.body_paragraph(doc, indent=not (previous == "equation" and payload[:1].islower()))
            K.add_rich_text(p, resolve_refs(payload, numbers))
        elif kind == "equation":
            img, width = next(eq_iter)
            K.equation_table(doc, img, width, numbers[payload["label"]])
        elif kind == "figure":
            K.picture(doc, HERE / "figures" / payload["file"], payload["width"])
            K.caption(doc, "Figure", numbers[payload["label"]], resolve_refs(payload["caption"], numbers))
        elif kind == "table":
            K.caption(doc, "Table", numbers[payload["label"]], resolve_refs(payload["caption"], numbers))
            header, *rows = [[resolve_refs(c, numbers) for c in r] for r in payload["rows"]]
            K.three_line_table(doc, header, rows, widths_cm=payload["widths"])
            K.body_paragraph(doc, indent=False)  # spacing after the table, as in the template
        elif kind == "refs":
            for entry in payload:
                K.reference_entry(doc, entry)
        previous = kind
    doc.core_properties.title = ""
    doc.core_properties.author = ""
    doc.save(str(out_docx))
    return numbers


def to_pdf(docx_path):
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent),
                    str(docx_path)], check=True, capture_output=True, timeout=600)
    return docx_path.with_suffix(".pdf")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--content", type=Path, default=HERE / "content.md")
    ap.add_argument("--summary", type=Path, default=ROOT / "results" / "summary.json")
    ap.add_argument("--out", type=Path, default=HERE / "Manuscript.docx")
    ap.add_argument("--no-pdf", action="store_true")
    args = ap.parse_args()
    numbers = build(args.content, args.out, args.summary)
    print(f"wrote {args.out} ({sum(1 for k in numbers if k.startswith('eq'))} equations, "
          f"{sum(1 for k in numbers if k.startswith('fig'))} figures, {sum(1 for k in numbers if k.startswith('tab'))} tables)")
    if not args.no_pdf:
        print("wrote", to_pdf(args.out))


if __name__ == "__main__":
    main()
