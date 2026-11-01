"""Helpers that write manuscript elements in the formats of the reference template.

Body paragraphs: Normal style, double spacing, first-line indent 420 twips, justified.
Display equations: borderless two-column table (7650 + 646 twips) with the equation centred and "(n)".
Tables: three-line rules (1.5 pt top/bottom, 0.75 pt under the header), 9 pt single-spaced cells.
Captions: centred paragraphs "Figure N. ..." below figures and "Table N. ..." above tables.
"""
from __future__ import annotations

import re
import subprocess
from pathlib import Path

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, Twips

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"

GREEK = {
    "alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "epsilon": "ε", "eta": "η", "theta": "θ",
    "kappa": "κ", "lambda": "λ", "mu": "μ", "nu": "ν", "xi": "ξ", "pi": "π", "rho": "ρ", "sigma": "σ",
    "tau": "τ", "phi": "φ", "varphi": "φ", "chi": "χ", "psi": "ψ", "omega": "ω",
    "Gamma": "Γ", "Delta": "Δ", "Theta": "Θ", "Lambda": "Λ", "Sigma": "Σ", "Phi": "Φ", "Omega": "Ω",
}
SYMBOLS = {
    "le": "≤", "leq": "≤", "ge": "≥", "geq": "≥", "approx": "≈", "sim": "~", "times": "×", "cdot": "·",
    "infty": "∞", "pm": "±", "to": "→", "propto": "∝", "ll": "≪", "gg": "≫", "neq": "≠", "simeq": "≃",
    "sum": "Σ", "partial": "∂", "ldots": "…", ",": " ", ";": " ", "quad": " ", " ": " ",
}
FUNCTIONS = {"ln", "log", "exp", "max", "min", "arg", "Var", "Re", "dB", "SNR"}
ACCENTS = {"bar": "̄", "hat": "̂", "tilde": "̃"}


# ----------------------------------------------------------------------------- paragraphs and runs

def set_spacing(paragraph, line=480, before=None, after=None):
    pf = paragraph.paragraph_format
    pf.line_spacing = line / 240.0
    if before is not None:
        pf.space_before = Twips(before)
    if after is not None:
        pf.space_after = Twips(after)


def body_paragraph(doc, indent=True, align="both"):
    p = doc.add_paragraph(style="Normal")
    set_spacing(p)
    if indent:
        p.paragraph_format.first_line_indent = Twips(420)
    p.alignment = {"both": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
    return p


def add_run(paragraph, text, italic=False, bold=False, sub=False, sup=False, size=None):
    run = paragraph.add_run(text)
    run.italic = italic or None
    run.bold = bold or None
    if sub:
        run.font.subscript = True
    if sup:
        run.font.superscript = True
    if size:
        run.font.size = Pt(size)
    return run


def _read_group(s, i):
    """Return (content, next_index) for a {...} group or a single token starting at s[i]."""
    if i >= len(s):
        return "", i
    if s[i] == "{":
        depth, j = 1, i + 1
        while j < len(s) and depth:
            depth += {"{": 1, "}": -1}.get(s[j], 0)
            j += 1
        return s[i + 1:j - 1], j
    if s[i] == "\\":
        m = re.match(r"\\([A-Za-z]+|.)", s[i:])
        return m.group(0), i + len(m.group(0))
    return s[i], i + 1


SPACED = set("=<>+") | {"≤", "≥", "≈", "×", "±", "→", "∝", "≠", "≃", "~"}


def math_runs(tex):
    """Convert a small LaTeX subset to (text, italic, vertical) runs; vertical is '', 'sub' or 'sup'.

    Relations and binary operators on the base line get surrounding spaces; a minus sign is binary only
    when it follows an operand.
    """
    out = []

    def last_char():
        text = "".join(t for t, _, v in out if v == "").rstrip()
        return text[-1] if text else ""

    def emit(text, italic, vert):
        if vert == "" and (text in SPACED or (text == "−" and last_char() not in "" "(=+−×/[{≤≥≈±")):
            text = f" {text} "
        if out and out[-1][1] == italic and out[-1][2] == vert:
            out[-1] = (out[-1][0] + text, italic, vert)
        else:
            out.append((text, italic, vert))

    def walk(s, vert="", roman=False):
        i = 0
        while i < len(s):
            c = s[i]
            if c in "_^":
                g, i = _read_group(s, i + 1)
                walk(g, "sub" if c == "_" else "sup", roman)
                continue
            if c == "\\":
                m = re.match(r"\\([A-Za-z]+|.)", s[i:])
                name = m.group(1)
                i += len(m.group(0))
                if name in GREEK:
                    emit(GREEK[name], not roman and name[0].islower(), vert)
                elif name in SYMBOLS:
                    emit(SYMBOLS[name], False, vert)
                elif name in ("mathrm", "text", "operatorname"):
                    g, i = _read_group(s, i)
                    walk(g, vert, roman=True)
                elif name in ACCENTS:
                    g, i = _read_group(s, i)
                    base = GREEK.get(g.strip("\\"), g)
                    emit(base + ACCENTS[name], not roman, vert)
                elif name == "frac":
                    a, i = _read_group(s, i)
                    b, i = _read_group(s, i)
                    walk(a, vert, roman)
                    emit("/", False, vert)
                    walk(f"({b})" if len(b) > 1 and not b.isalnum() else b, vert, roman)
                elif name == "sqrt":
                    g, i = _read_group(s, i)
                    emit("√", False, vert)
                    walk(g if len(g) <= 1 else f"({g})", vert, roman)
                elif name in FUNCTIONS:
                    nxt = s[i:].lstrip()[:1]
                    lead = " " if vert == "" and last_char().isalnum() else ""
                    trail = " " if vert == "" and (nxt.isalnum() or nxt == "\\") else ""
                    emit(lead + name + trail, False, vert)
                elif name in ("left", "right", "big", "Big", "bigl", "bigr"):
                    continue
                else:
                    emit(name, False, vert)
                continue
            if c == "{" or c == "}":
                i += 1
                continue
            if c.isalpha():
                m = re.match(r"[A-Za-z]+", s[i:])
                word = m.group(0)
                i += len(word)
                if word in FUNCTIONS or roman:
                    emit(word, False, vert)
                else:
                    emit(word, True, vert)
                continue
            if c == " ":
                i += 1
                continue
            emit({"-": "−", "*": "*", "'": "′"}.get(c, c), False, vert)
            i += 1

    walk(tex)
    return out


def add_rich_text(paragraph, text, size=None):
    """Add text with $...$ inline math, *italic* and **bold** spans."""
    tokens = re.split(r"(\$[^$]+\$|\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("$"):
            for t, italic, vert in math_runs(tok[1:-1]):
                add_run(paragraph, t, italic=italic, sub=vert == "sub", sup=vert == "sup", size=size)
        elif tok.startswith("**"):
            add_run(paragraph, tok[2:-2], bold=True, size=size)
        elif tok.startswith("*"):
            add_run(paragraph, tok[1:-1], italic=True, size=size)
        else:
            add_run(paragraph, tok, size=size)


# ----------------------------------------------------------------------------- display equations

EQ_PREAMBLE = r"""\documentclass[multi=eqn,border=1.5bp]{standalone}
\usepackage{amsmath}
\usepackage{unicode-math}
\setmainfont{Times New Roman}
\setmathfont{TeX Gyre Termes Math}
\newenvironment{eqn}{\fontsize{10.5bp}{13bp}\selectfont}{}
\begin{document}
"""


def render_equations(equations, outdir, dpi=600):
    """Typeset display equations with XeLaTeX (Times text, Termes math, 10.5 pt) and rasterise them.

    Returns a list of (png_path, width_cm) at natural size, one per equation.
    """
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    body = "\n".join(f"\\begin{{eqn}}$\\displaystyle {e}$\\end{{eqn}}" for e in equations)
    (outdir / "equations.tex").write_text(EQ_PREAMBLE + body + "\n\\end{document}\n", encoding="utf-8")
    run = subprocess.run(["xelatex", "-interaction=nonstopmode", "-halt-on-error", "equations.tex"],
                         cwd=outdir, capture_output=True, text=True)
    if run.returncode != 0:
        raise RuntimeError(run.stdout[-3000:])
    info = subprocess.run(["pdfinfo", "-f", "1", "-l", str(len(equations)), "equations.pdf"], cwd=outdir,
                          capture_output=True, text=True, check=True).stdout
    sizes = [float(w) for w, _ in re.findall(r"Page\s+\d+ size:\s+([\d.]+) x ([\d.]+) pts", info)]
    if len(sizes) != len(equations):
        raise RuntimeError(f"expected {len(equations)} equation pages, found {len(sizes)}")
    out = []
    for k, width_pt in enumerate(sizes, 1):
        stem = outdir / f"eq{k:02d}"
        subprocess.run(["pdftoppm", "-r", str(dpi), "-f", str(k), "-l", str(k), "-singlefile", "-png",
                        "equations.pdf", str(stem)], cwd=outdir, check=True)
        out.append((stem.with_suffix(".png"), width_pt / 72.0 * 2.54))
    return out


def no_grid_snap(paragraph):
    """Stop the template's line grid (docGrid pitch 312) from inflating image lines."""
    pPr = paragraph._p.get_or_add_pPr()
    snap = parse_xml(f'<w:snapToGrid {nsdecls("w")} w:val="0"/>')
    spacing = pPr.find(qn("w:spacing"))
    (spacing.addprevious if spacing is not None else pPr.append)(snap)


def _set_grid(table, widths_twips):
    for col, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths_twips):
        col.set(qn("w:w"), str(int(w)))


def _no_borders_xml():
    edges = "".join(f'<w:{e} w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
                    for e in ("top", "left", "bottom", "right", "insideH", "insideV"))
    return f"<w:tblBorders {nsdecls('w')}>{edges}</w:tblBorders>"


def equation_table(doc, image, width_cm, number):
    """Template equation layout: borderless 7650 + 646 twip table, equation centred, number at right."""
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    tblPr.append(parse_xml(_no_borders_xml()))
    tblPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
    _set_grid(table, (7650, 646))
    for cell, width in zip(table.rows[0].cells, (7650, 646)):
        cell.width = Twips(width)
        cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>'))
    eq_p, num_p = table.rows[0].cells[0].paragraphs[0], table.rows[0].cells[1].paragraphs[0]
    for p in (eq_p, num_p):
        set_spacing(p, line=240, before=120, after=120)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        no_grid_snap(p)
    eq_p.add_run().add_picture(str(image), width=Cm(width_cm))
    num_p.add_run(f"({number})")
    return table


# ----------------------------------------------------------------------------- data tables and figures

def caption(doc, label, number, text):
    p = body_paragraph(doc, indent=False, align="center")
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = label == "Table"
    add_run(p, f"{label} {number}. ")
    add_rich_text(p, text)
    return p


def _cell_borders(cell, top=None, bottom=None):
    parts = []
    if top:
        parts.append(f'<w:top w:val="single" w:sz="{top}" w:space="0" w:color="auto"/>')
    if bottom:
        parts.append(f'<w:bottom w:val="single" w:sz="{bottom}" w:space="0" w:color="auto"/>')
    if parts:
        cell._tc.get_or_add_tcPr().append(parse_xml(f"<w:tcBorders {nsdecls('w')}>{''.join(parts)}</w:tcBorders>"))


def three_line_table(doc, header, rows, widths_cm=None, font_pt=9):
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table._tbl.tblPr.append(parse_xml(_no_borders_xml()))
    if widths_cm:
        _set_grid(table, [w / 2.54 * 1440 for w in widths_cm])
    for r, values in enumerate([header] + rows):
        for c, value in enumerate(values):
            cell = table.rows[r].cells[c]
            if widths_cm:
                cell.width = Cm(widths_cm[c])
            cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>'))
            p = cell.paragraphs[0]
            set_spacing(p, line=240)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = r < len(rows)
            add_rich_text(p, str(value), size=font_pt)
            _cell_borders(cell, top=12 if r == 0 else None,
                          bottom=6 if r == 0 else (12 if r == len(rows) else None))
    return table


def picture(doc, path, width_cm):
    p = body_paragraph(doc, indent=False, align="center")
    p.paragraph_format.keep_with_next = True
    no_grid_snap(p)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    return p


def heading(doc, text, level, numbered=True):
    p = doc.add_paragraph(style=f"Heading {level}")
    set_spacing(p)
    if level == 1 and numbered:
        numPr = parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="2"/></w:numPr>')
        p._p.get_or_add_pPr().insert(0, numPr)
    add_rich_text(p, text)
    return p


def title(doc, text):
    p = doc.add_paragraph(style="Title")
    set_spacing(p)
    add_run(p, text)
    return p


def reference_entry(doc, text):
    p = body_paragraph(doc, indent=False, align="both")
    p.paragraph_format.left_indent = Twips(420)
    p.paragraph_format.first_line_indent = Twips(-420)
    add_rich_text(p, text)
    return p
